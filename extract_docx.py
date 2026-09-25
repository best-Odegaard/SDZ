from docx import Document
doc = Document('D:/geek/外/赛搭子——用户操作手册2.docx')
with open('D:/geek/外/doc_content.txt', 'w', encoding='utf-8') as f:
    f.write('=== PARAGRAPHS ===\n')
    for para in doc.paragraphs:
        if para.text.strip():
            f.write(para.text + '\n')
    f.write('\n=== TABLES ===\n')
    for i, table in enumerate(doc.tables):
        f.write(f'--- Table {i+1} ---\n')
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            f.write(' | '.join(cells) + '\n')
print('Done')
